import os
import sys
import subprocess

def install(package):
    subprocess.check_call([sys.executable, "-m", "pip", "install", package])

try:
    from fpdf import FPDF
except ImportError:
    install("fpdf")
    from fpdf import FPDF

try:
    from docx import Document
except ImportError:
    install("python-docx")
    from docx import Document

# Content of the proposal
title = "Proposal for Real Sample Data Collection\nSoil Microbiome Analysis for Mango Crops in Malihabad (U.P.)"
content = """
1. Introduction & Objectives
This proposal outlines the strategy for real sample data collection to support the research project "Hybrid AI with Explainable AI for Soil Microbiome Analysis and Predictive Modeling for Mango Crop at Malihabad (U.P.)". The objective is to gather high-quality, representative soil microbiome and physicochemical data to train and validate hybrid AI models for predicting mango crop yield and soil health.

2. Study Area and Plot Size
Location: Major mango-growing regions in Malihabad, Uttar Pradesh.
Orchard Selection: Three distinct representative orchards will be selected to capture variability in historical productivity (High, Medium, and Low yield histories) and soil management practices.
Plot Size: A standard plot size of 1 Acre to 1 Hectare per selected orchard.

3. Soil Testing Program

3.1 Sampling Frequency
Sampling will be conducted quarterly (4 times a year) to capture seasonal and phenological variations essential for the mango crop cycle:
- Q1 (Jan - Mar): Pre-flowering and flowering stage.
- Q2 (Apr - Jun): Fruit setting and development (Pre-monsoon).
- Q3 (Jul - Sep): Post-harvest and vegetative growth (Monsoon).
- Q4 (Oct - Dec): Dormancy/Preparation for the next cycle (Winter).

3.2 Parameters for Testing
The following parameters will be recorded for each sample:
- Physicochemical Parameters: Soil pH, Soil Moisture, Soil Temperature, Organic Carbon (OC), and Primary Macronutrients (N, P, K).
- Biological Parameters: Soil microbial DNA extraction for metagenomic sequencing (OTU clustering to identify microbial communities).
- Agronomic Metadata: Historical/current yield data, tree age/variety (e.g., Dasheri), and management practices (fertilizers, pesticides, irrigation types).

4. Sample Size and Collection Methodology

4.1 Collection Technique
Location: Samples will be collected from the rhizosphere, directly beneath the canopy drip line of the trees, at a depth of 15-30 cm where root activity is maximum.
Composite Sampling: Within each of the 3 designated orchards, a zig-zag walking pattern will be used across the plot. 5 to 10 sub-samples will be collected from different trees. These sub-samples will be thoroughly mixed to form one composite sample per orchard.

4.2 Sample Size Summary
Minimum Samples: 3 composite samples (1 per distinct orchard/zone) every 3 months.
Annual Total: 12 composite datasets per year.
Integration with IoT: Continuous environmental data logging (using the 10 proposed IoT sensors) will supplement these physical samples to provide a rich, unified data warehouse for Hybrid AI modeling.
"""

# Generate PDF
class PDF(FPDF):
    def header(self):
        self.set_font('Arial', 'B', 14)
        self.cell(0, 10, 'Proposal for Real Sample Data Collection', 0, 1, 'C')
        self.ln(5)

pdf = PDF()
pdf.add_page()
pdf.set_font('Arial', 'B', 12)
pdf.multi_cell(0, 8, title)
pdf.ln(5)
pdf.set_font('Arial', '', 11)

for line in content.strip().split('\n'):
    if line.startswith('1.') or line.startswith('2.') or line.startswith('3.') or line.startswith('4.'):
        pdf.set_font('Arial', 'B', 11)
        pdf.cell(0, 8, line, 0, 1)
        pdf.set_font('Arial', '', 11)
    else:
        # handle ascii issues
        line = line.encode('latin-1', 'replace').decode('latin-1')
        pdf.multi_cell(0, 6, line)

pdf_filename = 'Real_Sample_Data_Collection_Proposal.pdf'
pdf.output(pdf_filename)

# Generate DOCX
doc = Document()
doc.add_heading('Proposal for Real Sample Data Collection', 0)
doc.add_heading('Soil Microbiome Analysis for Mango Crops in Malihabad (U.P.)', level=1)

for paragraph in content.strip().split('\n\n'):
    if paragraph.strip():
        lines = paragraph.split('\n')
        # Check if the first line is a main heading
        if lines[0].startswith('1.') or lines[0].startswith('2.') or lines[0].startswith('3.') or lines[0].startswith('4.'):
            doc.add_heading(lines[0], level=2)
            for line in lines[1:]:
                if line.strip():
                    doc.add_paragraph(line)
        else:
            for line in lines:
                 if line.strip():
                     doc.add_paragraph(line)

docx_filename = 'Real_Sample_Data_Collection_Proposal.docx'
doc.save(docx_filename)

print(f"Successfully created {pdf_filename} and {docx_filename}")
